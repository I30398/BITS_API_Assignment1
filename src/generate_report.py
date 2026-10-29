"""Build a review or submission report from current implementation evidence."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

from src import config
from src.artifact_io import source_fingerprint


def read_json(path, default=None):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def resource_link(doc, label, url):
    """Keep submitted resource URLs visible and clickable in Word and PDF."""
    paragraph = doc.add_paragraph(label + ": ")
    paragraph.paragraph_format.space_after = Pt(5)
    for label_run in paragraph.runs:
        label_run.font.size = Pt(10)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    props.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def table(doc, columns, rows, widths=None):
    result = doc.add_table(rows=1, cols=len(columns))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    if widths is None:
        widths = [6.5 / len(columns)] * len(columns)
    for col, width in zip(result.columns, widths):
        col.width = Inches(width)
    for cell, title in zip(result.rows[0].cells, columns):
        cell.text = title
    for row in rows:
        for cell, value in zip(result.add_row().cells, row):
            cell.text = str(value)
    repeat = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(result.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                element = OxmlElement(f"w:{edge}")
                for key, val in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    element.set(qn(f"w:{key}"), val)
                borders.append(element)
            props.append(borders)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "243746" if row_index == 0 else
                        ("F2F5F7" if row_index % 2 else "FFFFFF"))
            props.append(shading)
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:w"), "85")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            props.append(margins)
            for paragraph in cell.paragraphs:
                if (row_index < len(result.rows) - 1 and
                        (len(result.rows) <= 9 or row_index <= 1)):
                    paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return result


def image(doc, path, caption, width=6.2):
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_paragraph.runs:
        run.italic = True
        run.font.size = Pt(9)


def validate_group(group):
    import re
    members = group.get("members", [])
    if not group.get("group_id") or not 2 <= len(members) <= 6:
        raise ValueError("Provide group_id and 2 to 6 actual members in config/group.json.")
    if not re.fullmatch(r"[A-Za-z0-9_-]+", group["group_id"]):
        raise ValueError("Group ID must contain letters, digits, underscores or hyphens.")
    for member in members:
        if not all(member.get(k) for k in ["bits_id", "name", "contribution"]):
            raise ValueError("Every member needs a BITS ID, name and qualitative contribution.")
    if abs(sum(float(m.get("percentage", 0)) for m in members) - 100) > 0.001:
        raise ValueError("Member contribution percentages must total 100.")
    return group


def validate_snapshot(root):
    """Verify the frozen output set without modifying its original manifests."""
    root = Path(root).resolve()
    inventory = read_json(root / "inventory.json")
    if not inventory:
        raise ValueError("The frozen snapshot needs an inventory.")
    for item in inventory:
        path = (root / item["path"]).resolve()
        if not path.is_relative_to(root) or not path.is_file():
            raise ValueError("Invalid snapshot path: " + item["path"])
        if hashlib.sha256(path.read_bytes()).hexdigest() != item["sha256"]:
            raise ValueError("Snapshot checksum mismatch: " + item["path"])
    hosting = read_json(root / "hosting.json", {})
    verification = read_json(root / "outputs/verification.json", {})
    if (hosting.get("provider") != "BITS Prayogshala" or not hosting.get("hostname")
            or not hosting.get("functional_checks_passed")
            or len(verification.get("checks", {})) < 14
            or not all(verification["checks"].values())
            or len(verification.get("scheduled_parent_runs", [])) < 2):
        raise ValueError("The snapshot lacks successful hosting and scheduled-run evidence.")
    screenshots = root / "screenshots"
    captures = read_json(screenshots / "capture_inventory.json", {})
    # The capture record has a different host and time from the frozen run.
    # Validate its own hashes rather than implying it belongs to that same run.
    names = set()
    for item in captures.get("captures", []):
        path = (screenshots / item["file"]).resolve()
        if (not path.is_relative_to(screenshots) or not path.is_file()
                or hashlib.sha256(path.read_bytes()).hexdigest() != item["sha256"]):
            raise ValueError("Screenshot checksum mismatch: " + item["file"])
        names.add(item["file"])
    required = {"01_dashboard.png", "02_flow_runs.png", "03_deployments.png",
                "05_ml_run_logs.png", "06_data_run_logs.png", "07_hosting.png"}
    if not required <= names or not captures.get("vm_host"):
        raise ValueError("The cloud snapshot needs all six documented captures.")
    hosting = dict(hosting, screenshot_host=captures["vm_host"])
    return hosting, verification


def build(draft=False, group_path=None, snapshot_root=None, output_dir=None):
    group_path = group_path or config.PROJECT_ROOT / "config/group.json"
    group = read_json(group_path, {})
    hosting = verification = None
    if snapshot_root:
        snapshot_root = Path(snapshot_root).resolve()
        hosting, verification = validate_snapshot(snapshot_root)
    outputs = snapshot_root / "outputs" if snapshot_root else config.OUTPUTS_DIR
    models, eda = outputs / "models", outputs / "eda"
    screenshots = snapshot_root / "screenshots" if snapshot_root else config.SCREENSHOTS_DIR
    metrics = read_json(models / "latest_metrics.json")
    run = read_json(outputs / "latest_run.json", {})
    api = read_json(outputs / "api_details.json", {})
    if not metrics:
        raise ValueError("Run the revised workflow before generating a report.")
    if metrics.get("source_sha256") != source_fingerprint():
        raise ValueError("Current evidence was produced by different source code. Rerun the workflow.")
    if not draft:
        validate_group(group)
        if not group.get("video_url"):
            raise ValueError("Add the uploaded demonstration video's URL before final generation.")
        if not group.get("cloud_evidence_confirmed"):
            raise ValueError("Cloud hosting evidence is not confirmed. Use --draft for review.")
        for name in ["01_dashboard.png", "02_flow_runs.png", "03_deployments.png",
                     "05_ml_run_logs.png", "06_data_run_logs.png"]:
            if not (screenshots / name).exists():
                raise ValueError(f"Capture current dashboard evidence: {name}")
        if api.get("source_sha256") != source_fingerprint():
            raise ValueError("Capture fresh API details from the current implementation.")
    doc = Document()
    # The bundled Word template may carry a blue rule beneath its Title style.
    for root in [doc.styles.element, doc.element]:
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(1)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ["Heading 1", "Heading 2"]:
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.add_paragraph("Telecom Customer Churn Prediction", "Title")
    doc.add_paragraph("CCZG506 API driven Cloud Native Solutions", "Subtitle")
    doc.add_paragraph("Assignment I" + (" review copy" if draft else ""))
    doc.add_paragraph(
        "This project identifies customers at risk of leaving a telecom service. "
        "A scheduled Prefect workflow performs data preparation and exploration, "
        "then trains and evaluates two classifiers. Prefect APIs expose the "
        "application's flows, deployments and execution details.")
    doc.add_paragraph(f"Group number: {group.get('group_id', '________________')}")
    members = group.get("members") or [{"bits_id": "", "name": "", "contribution": "", "percentage": ""} for _ in range(2)]
    table(doc, ["BITS ID", "Name", "Contribution", "%"],
          [[m.get(k, "") for k in ["bits_id", "name", "contribution", "percentage"]] for m in members],
          [1.1, 1.3, 3.6, 0.5])
    doc.add_paragraph("Execution and evidence", "Heading 1")
    doc.add_paragraph(f"Execution environment: {run.get('environment', 'local')}. "
        f"Latest completed cycle: {run.get('completed_at', 'not captured')}. "
        "The original repository outputs are retained separately from these results.")
    if hosting:
        doc.add_paragraph(
            f"The application ran on the {hosting['provider']} CCZG506 virtual lab, "
            f"host {hosting['hostname']}. The frozen output set was captured at "
            f"{hosting['captured_at']}. Its inventory hashes were verified after transfer.")
        doc.add_paragraph(
            f"All {len(verification['checks'])} functional checks passed. "
            f"The API recorded {len(verification['scheduled_parent_runs'])} completed "
            "scheduled parent cycles with a 120-second interval and concurrency of one. "
            "Provider and dashboard screenshots were captured later, on 26 October 2026, "
            f"after services resumed on replacement host {hosting['screenshot_host']}. "
            "The screenshots and frozen results retain their separate capture times.")
        doc.add_paragraph(
            "The verification manifest retains its default local environment label and "
            "does not independently certify hosting. Cloud execution is supported by the "
            "VM metadata, provider status and application captures together. "
            "The dashboard is private and session-bound; the lab's inactivity limits apply.")
    elif not group.get("cloud_evidence_confirmed"):
        doc.add_paragraph("The evidence in this review demonstrates local execution. "
                          "Cloud hosting has not yet been demonstrated.")
    doc.add_paragraph("1 Data pipeline", "Heading 1")
    doc.add_paragraph("1.1 Business understanding and data ingestion", "Heading 2")
    doc.add_paragraph("The business objective is to support retention decisions by predicting "
        "customer churn from account, service and demographic attributes. The public IBM "
        "Telco Customer Churn dataset contains 7043 records and 21 columns. Churn is a "
        "binary Yes or No target; customerID is excluded from model features.")
    doc.add_paragraph(f"Dataset source: {config.DATASET_URL}")
    image(doc, eda / "univariate_churn_distribution.png", "Figure 1 Customer churn class balance", 4.5)
    doc.add_paragraph("1.2 Preprocessing", "Heading 2")
    doc.add_paragraph("The workflow coerces TotalCharges to numeric, identifies missing values, "
        "and induces 352 missing entries in each of tenure, MonthlyCharges and TotalCharges "
        "using a fixed random seed. Full-data EDA uses median imputation and Min-Max "
        "normalisation. The pre-imputation snapshot is passed separately to ML so that "
        "test records cannot influence training preprocessing.")
    stats = pd.read_csv(outputs / "summary_statistics.csv", index_col=0)
    table(doc, ["Feature", "Count", "Mean", "Std", "Min", "Max"],
          [[idx, int(row['count']), *[f"{row[k]:.2f}" for k in ['mean', 'std', 'min', 'max']]]
           for idx, row in stats.iterrows()], [1.45, 0.7, 1.1, 1.1, 1.0, 1.15])
    types = pd.read_csv(outputs / "data_types.csv")
    table(doc, ["Column", "Data type"], types.values.tolist(), [4.3, 2.2])
    preview = pd.read_csv(outputs / "normalisation_preview.csv")
    table(doc, ["Numeric feature", "Original example", "Normalised example"],
          [[c, f"{preview.iloc[0][c]:.2f}", f"{preview.iloc[0][c + '_norm']:.4f}"]
           for c in config.NUMERIC_COLUMNS], [2.5, 2, 2])
    doc.add_paragraph("1.3 Exploratory data analysis", "Heading 2")
    doc.add_paragraph("Tenure is binned into five bands. Encoding mappings and bin counts "
        "are saved. Pearson coefficients are calculated for numeric features and the binary "
        "target. Cramer's V measures categorical association without assuming an ordering "
        "between category codes. Random Forest exploratory importance is aggregated "
        "from one-hot category levels, using each original numeric feature once.")
    assoc = pd.read_csv(outputs / "categorical_associations.csv").head(6)
    table(doc, ["Categorical feature", "Association with churn"],
          [[row.feature, f"{row.cramers_v:.3f}"] for row in assoc.itertuples()], [3.5, 3])
    for name, caption in [
        ("correlation_heatmap.png", "Figure 2 Numeric correlations"),
        ("feature_importance.png", "Figure 3 Exploratory feature importance"),
        ("univariate_tenure_histogram.png", "Figure 4 Univariate tenure distribution"),
        ("bivariate_churn_by_contract.png", "Figure 5 Churn rate by contract"),
        ("bivariate_charges_by_churn.png", "Figure 6 Charges by churn outcome")]:
        image(doc, eda / name, caption, 5.5)
    doc.add_paragraph("1.4 DataOps scheduling and monitoring", "Heading 2")
    doc.add_paragraph("One deployment runs every 120 seconds with concurrency limited to "
        "one cycle. DataOps completes before MLOps starts. Complete files are published by "
        "atomic replacement, and cross-process locks protect appended logs. Every task "
        "records its activity in the Prefect dashboard and JSONL logs; result tables are "
        "also available as Prefect artifacts.")
    for name, caption in [("03_deployments.png", "Figure 7 Deployment interval and status"),
                          ("02_flow_runs.png", "Figure 8 Scheduled execution history"),
                          ("06_data_run_logs.png", "Figure 9 Data preparation activity logs")]:
        image(doc, screenshots / name, caption)
    doc.add_paragraph("2 Machine learning pipeline", "Heading 1")
    doc.add_paragraph("2.1 Preparation training and evaluation", "Heading 2")
    doc.add_paragraph("Logistic Regression provides a linear baseline; Random Forest "
        "captures non-linear relationships. A stratified 70/30 split produces 4930 "
        "training records and 2113 test records. Each model contains its own fitted "
        "preprocessing pipeline: numeric median imputation and standardisation, plus "
        "categorical imputation and one-hot encoding. All fitted parameters are learned "
        "from training records. Test records are transformed without refitting.")
    table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC AUC"],
          [[m['model'], *[f"{m[k]:.4f}" for k in ['accuracy', 'precision', 'recall', 'f1_score', 'roc_auc']]]
           for m in metrics['models']], [1.6, .98, .98, .98, .98, .98])
    image(doc, models / "model_comparison.png", "Figure 10 Five metrics for each model")
    image(doc, models / "confusion_matrices.png", "Figure 11 Holdout confusion matrices")
    doc.add_paragraph("2.2 MLOps and interpretation", "Heading 2")
    doc.add_paragraph("Every cycle records accuracy, precision, recall, F1 and ROC AUC "
        "for both models. The metrics history records timestamps, pipeline version and "
        "source fingerprint. Repeating evaluation on this fixed dataset demonstrates "
        "monitoring and logging; it does not establish drift detection or performance "
        "on a changing production population. Recall measures the proportion of actual "
        "churners identified, while precision measures how many flagged customers churn.")
    best = metrics['best']
    doc.add_paragraph(
        f"{best['model']} performed better in this holdout comparison, with ROC AUC "
        f"{best['roc_auc']:.4f} and recall {best['recall']:.4f}. At the default "
        f"decision threshold, {100 * (1 - best['recall']):.1f}% of actual churners "
        "were missed. "
        "Retention use would require threshold selection and validation on later data.")
    image(doc, screenshots / "05_ml_run_logs.png", "Figure 12 ML task logs and metrics")
    doc.add_paragraph("3 API access", "Heading 1")
    doc.add_paragraph("The official Prefect client retrieves details through built-in APIs. "
        "Project filters keep unrelated flows out of the evidence. The saved response "
        "contains registered flows, deployment schedules and concurrency, recent run "
        "states, start times and durations. Aggregate statistics describe a bounded "
        "sample of up to 200 project runs and include child flows.")
    table(doc, ["Detail", "Retrieved value"], [
        ["Retrieved at", api.get('retrieved_at', 'Not captured')],
        ["Project flows", ', '.join(f['name'] for f in api.get('flows', []))],
        ["Deployments", ', '.join(d['name'] for d in api.get('deployments', []))],
        ["Recent executed runs", len(api.get('recent_flow_runs', []))],
        ["Sampled run states", str(api.get('run_statistics', {}).get('runs_by_state', {}))],
    ], [1.7, 4.8])
    image(doc, screenshots / "01_dashboard.png", "Figure 13 Dashboard overview")
    doc.add_paragraph("4 Reproduction and demonstration", "Heading 1")
    if hosting:
        doc.add_paragraph(
            "On the provisioned BITS VM, run .venv/bin/python scripts/lab_start.py "
            "from ~/workspace/BITS_API_Assignment1. The launcher restores the dedicated "
            "hard NFS mount, checks the persistent PostgreSQL database, then starts "
            "Prefect on 127.0.0.1:4207 and the two-minute schedule. All runtime and "
            "project files are under ~/workspace. Stop the launcher with Ctrl+C. "
            "See docs/DEPLOYMENT.md for setup, storage and private-access details.")
        image(doc, screenshots / "07_hosting.png", "Figure 14 BITS virtual lab hosting status", 4.2)
    else:
        doc.add_paragraph("Follow README.md for the tested Python 3.12 setup, local server "
            "and scheduled workflow. The local dashboard is at http://127.0.0.1:4207. "
            "Run python -m src.api_access to capture current API details. The Docker Compose "
            "package provides the same server and runner arrangement for local or cloud VM "
            "execution; selecting cloud configuration alone is not proof of deployment.")
    if group.get("video_url"):
        resource_link(doc, "Demonstration video", group["video_url"])
    if group.get("artifacts_url"):
        resource_link(doc, "Assignment artifacts", group["artifacts_url"])
    resource_link(doc, "Source code", "https://github.com/I30398/BITS_API_Assignment1")
    doc.add_paragraph("Dataset SHA256: " + metrics['raw_dataset_sha256'])
    output_dir = Path(output_dir) if output_dir else config.PROJECT_ROOT / "report/current"
    output_dir.mkdir(parents=True, exist_ok=True)
    name = "CCZG506_Assignment1_Review.docx" if draft else group['group_id'] + '.docx'
    output = output_dir / name
    doc.save(output)
    print(output)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--draft", action="store_true", help="Generate a review copy with incomplete submission details.")
    parser.add_argument("--group", type=Path)
    parser.add_argument("--snapshot-root", type=Path, help="Read a verified frozen cloud snapshot.")
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    build(args.draft, args.group, args.snapshot_root, args.output_dir)
